import yfinance as yf
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BG_DARK   = "0D0D14"
BG_HEADER = "1A1A2E"
FG_PINK   = "E879F9"
FG_GRAY   = "888888"
GREEN     = "00B050"

def _hfill(color): return PatternFill("solid", fgColor=color)
def _hfont(color, bold=True, size=11): return Font(color=color, bold=bold, size=size)

def _row(ws, r, label, value=None, fmt=None, formula=None, bold=False, color=None):
    lc = ws.cell(row=r, column=1, value=label)
    lc.font = Font(bold=bold)
    vc = ws.cell(row=r, column=2, value=formula if formula else value)
    if fmt:
        vc.number_format = fmt
    if color:
        vc.font = Font(bold=True, color=color)
    return f"B{r}"

def build_dcf_sheet(ws, ticker):
    stock = yf.Ticker(ticker)
    cf    = stock.cashflow
    bs    = stock.balance_sheet
    inc   = stock.financials
    info  = stock.info

    op_cf      = float(cf.loc["Operating Cash Flow"].iloc[0]) / 1e6
    capex      = float(cf.loc["Capital Expenditure"].iloc[0]) / 1e6
    fcf        = op_cf + capex
    total_debt = float(bs.loc["Total Debt"].iloc[0]) / 1e6
    cash       = float(bs.loc["Cash And Cash Equivalents"].iloc[0]) / 1e6
    shares     = float(info.get("sharesOutstanding", 0)) / 1e6
    cur_price  = float(info.get("currentPrice", 0))
    beta       = float(info.get("beta", 1.0))
    rf         = 0.043
    erp        = 0.055
    terminal_g = 0.03
    years      = 5

    revs = inc.loc["Total Revenue"].dropna()
    rev_growth = float((revs.iloc[0] / revs.iloc[1]) - 1) if len(revs) >= 2 else 0.10

    r = 1
    title = ws.cell(row=r, column=1, value=f"{ticker}  —  DCF Valuation Model  ($ Millions)")
    title.font = Font(bold=True, size=14, color=FG_PINK)
    title.fill = _hfill(BG_DARK)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    r += 2

    def sec_header(label):
        nonlocal r
        c = ws.cell(row=r, column=1, value=label)
        c.fill = _hfill(BG_HEADER)
        c.font = _hfont(FG_PINK)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        r += 1

    # SECTION A: MARKET DATA
    sec_header("A  >  MARKET & COMPANY DATA  (auto-pulled)")

    row_price = r;  _row(ws, r, "Current Share Price ($)",       cur_price,  '$#,##0.00'); r+=1
    row_shares= r;  _row(ws, r, "Shares Outstanding (M)",        shares,     '#,##0.00');  r+=1
    row_mktcap= r;  _row(ws, r, "Market Capitalisation ($M)",    None,       '$#,##0.00', formula=f"=B{row_price}*B{row_shares}"); r+=1
    row_beta  = r;  _row(ws, r, "Beta",                          beta,       '0.00');      r+=1
    row_debt  = r;  _row(ws, r, "Total Debt ($M)",               total_debt, '$#,##0.00'); r+=1
    row_cash  = r;  _row(ws, r, "Cash & Equivalents ($M)",       cash,       '$#,##0.00'); r+=1
    row_netd  = r;  _row(ws, r, "Net Debt ($M)",                 None,       '$#,##0.00', formula=f"=B{row_debt}-B{row_cash}"); r+=1
    row_opcf  = r;  _row(ws, r, "Operating Cash Flow ($M)",      op_cf,      '$#,##0.00'); r+=1
    row_capex = r;  _row(ws, r, "Capital Expenditure ($M)",      capex,      '$#,##0.00'); r+=1
    row_fcf   = r;  _row(ws, r, "Free Cash Flow ($M)",           None,       '$#,##0.00', formula=f"=B{row_opcf}+B{row_capex}"); r+=1
    r += 1

    # SECTION B: WACC
    sec_header("B  >  WACC  (edit yellow cells)")

    rRF  = r; _row(ws, r, "Risk-Free Rate (10-yr Treasury)",  rf,    '0.00%'); r+=1
    rERP = r; _row(ws, r, "Equity Risk Premium",              erp,   '0.00%'); r+=1
    rB   = r; _row(ws, r, "Beta (from market data)",          None,  '0.00',  formula=f"=B{row_beta}"); r+=1
    rCE  = r; _row(ws, r, "Cost of Equity (CAPM)",            None,  '0.00%', formula=f"=B{rRF}+B{rB}*B{rERP}"); r+=1
    rKD  = r; _row(ws, r, "Pre-tax Cost of Debt",             0.05,  '0.00%'); r+=1
    rTX  = r; _row(ws, r, "Tax Rate",                         0.21,  '0.00%'); r+=1
    rAD  = r; _row(ws, r, "After-tax Cost of Debt",           None,  '0.00%', formula=f"=B{rKD}*(1-B{rTX})"); r+=1
    rEW  = r; _row(ws, r, "Equity Weight (E/V)",              None,  '0.00%', formula=f"=B{row_mktcap}/(B{row_mktcap}+B{row_debt})"); r+=1
    rDW  = r; _row(ws, r, "Debt Weight (D/V)",                None,  '0.00%', formula=f"=1-B{rEW}"); r+=1
    rWACC= r; _row(ws, r, "WACC",                             None,  '0.00%', formula=f"=B{rEW}*B{rCE}+B{rDW}*B{rAD}", bold=True, color=FG_PINK); r+=1
    r += 1

    # SECTION C: GROWTH
    sec_header("C  >  GROWTH ASSUMPTIONS  (edit yellow cells)")

    rFCF = r; _row(ws, r, "Base FCF ($M)",                    fcf,       '$#,##0.00', formula=f"=B{row_fcf}"); r+=1
    rGR  = r; _row(ws, r, "FCF Growth Rate (Yrs 1-5)",        rev_growth,'0.00%'); r+=1
    rTG  = r; _row(ws, r, "Terminal Growth Rate",             terminal_g,'0.00%'); r+=1
    rYR  = r; _row(ws, r, "Projection Years",                 years,     '0'); r+=1
    r += 1

    # SECTION D: PROJECTED FCFs
    sec_header("D  >  PROJECTED FREE CASH FLOWS")

    hdr_fill = _hfill(BG_HEADER)
    hdr_font = _hfont(FG_PINK)
    ws.cell(row=r, column=1, value="").fill = hdr_fill
    for y in range(1, years+1):
        c = ws.cell(row=r, column=y+1, value=f"Year {y}")
        c.fill = hdr_fill; c.font = hdr_font
        c.alignment = Alignment(horizontal="center")
    r += 1

    ws.cell(row=r, column=1, value="Projected FCF ($M)").font = Font(bold=True)
    proj_cells = []
    for y in range(1, years+1):
        col_l = get_column_letter(y+1)
        c = ws.cell(row=r, column=y+1, value=f"=B{rFCF}*(1+B{rGR})^{y}")
        c.number_format = '$#,##0.00'
        proj_cells.append(f"{col_l}{r}")
    rPROJ = r; r += 1

    ws.cell(row=r, column=1, value="Discounted FCF ($M)").font = Font(bold=True)
    disc_cells = []
    for y in range(1, years+1):
        col_l = get_column_letter(y+1)
        c = ws.cell(row=r, column=y+1, value=f"={col_l}{rPROJ}/(1+B{rWACC})^{y}")
        c.number_format = '$#,##0.00'
        disc_cells.append(f"{col_l}{r}")
    rDISC = r; r += 1
    r += 1

    # SECTION E: VALUATION
    sec_header("E  >  VALUATION SUMMARY")

    last_col = get_column_letter(years+1)

    rSPV = r; _row(ws, r, "Sum of PV of FCFs ($M)",        None, '$#,##0.00', formula=f"=SUM({disc_cells[0]}:{disc_cells[-1]})"); r+=1
    rTV  = r; _row(ws, r, "Terminal Value ($M)",            None, '$#,##0.00', formula=f"=({last_col}{rPROJ}*(1+B{rTG}))/(B{rWACC}-B{rTG})"); r+=1
    rPTV = r; _row(ws, r, "PV of Terminal Value ($M)",      None, '$#,##0.00', formula=f"=B{rTV}/(1+B{rWACC})^B{rYR}"); r+=1
    rEV  = r; _row(ws, r, "Enterprise Value ($M)",          None, '$#,##0.00', formula=f"=B{rSPV}+B{rPTV}", bold=True); r+=1
    rD2  = r; _row(ws, r, "(-) Total Debt ($M)",            None, '$#,##0.00', formula=f"=B{row_debt}"); r+=1
    rC2  = r; _row(ws, r, "(+) Cash ($M)",                  None, '$#,##0.00', formula=f"=B{row_cash}"); r+=1
    rEQV = r; _row(ws, r, "Equity Value ($M)",              None, '$#,##0.00', formula=f"=B{rEV}-B{rD2}+B{rC2}", bold=True); r+=1
    rSH2 = r; _row(ws, r, "Shares Outstanding (M)",         None, '#,##0.00',  formula=f"=B{row_shares}"); r+=1
    rIV  = r; _row(ws, r, "Intrinsic Value Per Share ($)",  None, '$#,##0.00', formula=f"=B{rEQV}/B{rSH2}", bold=True, color=FG_PINK); r+=1
    rCP  = r; _row(ws, r, "Current Market Price ($)",       None, '$#,##0.00', formula=f"=B{row_price}"); r+=1
    rUP  = r; _row(ws, r, "Upside / (Downside)",            None, '0.00%',     formula=f"=(B{rIV}-B{rCP})/B{rCP}", bold=True, color=GREEN); r+=1

    # Column widths
    ws.column_dimensions['A'].width = 38
    ws.column_dimensions['B'].width = 16
    for y in range(1, years+1):
        ws.column_dimensions[get_column_letter(y+2)].width = 14

    # Highlight editable cells yellow
    yellow = PatternFill("solid", fgColor="FFFF00")
    for er in [rRF, rERP, rKD, rTX, rGR, rTG]:
        ws[f"B{er}"].fill = yellow
        ws[f"B{er}"].font = Font(color="000000", bold=True)


def format_sheet(ws, df, title):
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14, color=FG_PINK)
    ws.cell(row=2, column=1, value="$ Millions").font = Font(italic=True, color=FG_GRAY)

    hdr_fill = _hfill(BG_HEADER)
    hdr_font = _hfont(FG_PINK)
    for c_idx, col in enumerate(df.columns, start=2):
        cell = ws.cell(row=2, column=c_idx, value=str(col)[:10])
        cell.fill = hdr_fill; cell.font = hdr_font
        cell.alignment = Alignment(horizontal="center")

    for r_idx, (index, row) in enumerate(df.iterrows(), start=3):
        ws.cell(row=r_idx, column=1, value=index).font = Font(bold=True)
        for c_idx, val in enumerate(row, start=2):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            if pd.notna(val) and isinstance(val, (int, float)):
                if abs(val) < 1000:
                    cell.value = round(val, 4)
                    cell.number_format = '#,##0.0000'
                else:
                    cell.value = round(val / 1_000_000, 2)
                    cell.number_format = '$#,##0.00'
            cell.alignment = Alignment(horizontal="right")

    for col in ws.columns:
        max_len = max((len(str(cell.value)) for cell in col if cell.value), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 32)


def export_to_excel(ticker):
    stock = yf.Ticker(ticker)
    wb = Workbook()

    ws_dcf = wb.active
    ws_dcf.title = "DCF Model"
    build_dcf_sheet(ws_dcf, ticker)

    ws1 = wb.create_sheet("Income Statement")
    format_sheet(ws1, stock.financials, "Income Statement")

    ws2 = wb.create_sheet("Balance Sheet")
    format_sheet(ws2, stock.balance_sheet, "Balance Sheet")

    ws3 = wb.create_sheet("Cash Flow")
    format_sheet(ws3, stock.cashflow, "Cash Flow")

    filename = f"{ticker}_DCF.xlsx"
    wb.save(filename)
    print(f"\n  Saved: {filename}")
    print(f"  DCF Model: fully formula-driven | yellow cells = editable assumptions")

if __name__ == "__main__":
    ticker = input("Enter ticker: ").upper()
    export_to_excel(ticker)
    